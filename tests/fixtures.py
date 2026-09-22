"""Generated fixture documents (obviously fictional Spanish content) for the tests and the selftest."""

from __future__ import annotations

from pathlib import Path

FUNES = (
    "Ireneo Funes, el memorioso de la aldea imaginaria de Valdeniebla, recordaba cada hoja de cada árbol de cada monte. "
    "Su memoria era como un vertedero de basuras: nada se perdía, nada se ordenaba. "
    "Pensar es olvidar diferencias, generalizar, abstraer; en el mundo abarrotado de Funes no había sino detalles."
)
JARDIN = (
    "El jardín de senderos que se bifurcan es una imagen incompleta, pero no falsa, del universo tal como lo concebía Ts'ui Pên. "
    "Creía en infinitas series de tiempos, en una red creciente y vertiginosa de tiempos divergentes, convergentes y paralelos."
)
BIBLIOTECA = (
    "La biblioteca de Valdeniebla es total: sus anaqueles registran todas las posibles combinaciones de los veintitantos símbolos ortográficos. "
    "Cuando se proclamó que la biblioteca abarcaba todos los libros, la primera impresión fue de extravagante felicidad."
)
BUDGET_EN = (
    "The quarterly maintenance budget for the fictional Valdeniebla observatory was approved by the committee without amendments. "
    "Spending on telescope mirrors will be reviewed in the next fiscal year."
)


def make_pdf(path: Path, pages: list[str] | None = None) -> Path:
    import pymupdf as fitz

    pages = pages or [FUNES, JARDIN, BIBLIOTECA]
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 780), text, fontsize=11, fontname="helv")
    doc.set_metadata({"title": path.stem.replace("_", " ").title()})
    doc.save(str(path))
    doc.close()
    return path


def make_scanned_pdf(path: Path, pages: int = 2) -> Path:
    """A PDF whose pages carry only an image: no text layer at all."""
    import pymupdf as fitz

    doc = fitz.open()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 40), False)
    pix.clear_with(200)
    for _ in range(pages):
        page = doc.new_page()
        page.insert_image(fitz.Rect(100, 100, 300, 300), pixmap=pix)
    doc.save(str(path))
    doc.close()
    return path


def make_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_heading("Memoria del máster", level=0)
    document.add_heading("Introducción", level=1)
    document.add_paragraph(FUNES)
    document.add_heading("Método", level=1)
    document.add_paragraph(JARDIN)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Variable"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "senderos"
    table.cell(1, 1).text = "infinitos"
    document.add_heading("Conclusiones", level=1)
    document.add_paragraph(BIBLIOTECA)
    document.save(str(path))
    return path


def make_md(path: Path) -> Path:
    path.write_text(
        "# Apuntes de lectura\n\nNotas sueltas sobre memoria y olvido.\n\n## Funes\n\n" + FUNES + "\n\n## El jardín\n\n" + JARDIN +
        "\n\n```\n# esto no es un título\n```\n\n## Biblioteca\n\n" + BIBLIOTECA + "\n",
        encoding="utf-8",
    )
    return path


def make_txt(path: Path, text: str = BUDGET_EN) -> Path:
    path.write_text(text + "\n", encoding="utf-8")
    return path


def make_html(path: Path) -> Path:
    path.write_text(
        "<html><head><title>Página de pruebas</title><style>p{color:red}</style></head><body>"
        "<h1>Ensayo sobre la memoria</h1><p>" + FUNES + "</p><script>var x = 1;</script>"
        "<h2>Los senderos</h2><p>" + JARDIN + "</p></body></html>",
        encoding="utf-8",
    )
    return path


def make_csv(path: Path) -> Path:
    path.write_text("titulo,autor,anyo\nEl memorioso,Anónimo de Valdeniebla,1942\nSenderos,Anónimo de Valdeniebla,1941\n", encoding="utf-8")
    return path


def make_epub(path: Path) -> Path:
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("urn:uuid:fixture-0001")
    book.set_title("Ficciones de Valdeniebla")
    book.set_language("es")
    book.add_author("Anónimo de Valdeniebla")
    chapters = []
    for index, (title, body) in enumerate([("Funes el memorioso", FUNES), ("El jardín", JARDIN), ("La biblioteca", BIBLIOTECA)], start=1):
        chapter = epub.EpubHtml(title=title, file_name=f"cap{index}.xhtml", lang="es")
        chapter.content = f"<html><body><h1>{title}</h1><p>{body}</p></body></html>"
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *chapters]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        epub.write_epub(str(path), book)
    return path


def make_library(folder: Path) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "sub").mkdir(exist_ok=True)
    return {
        "pdf": make_pdf(folder / "cuentos_valdeniebla.pdf"),
        "scanned": make_scanned_pdf(folder / "escaneado.pdf"),
        "docx": make_docx(folder / "memoria_master.docx"),
        "md": make_md(folder / "apuntes.md"),
        "txt": make_txt(folder / "sub" / "budget.txt"),
        "html": make_html(folder / "ensayo.html"),
        "csv": make_csv(folder / "catalogo.csv"),
        "epub": make_epub(folder / "ficciones.epub"),
    }
